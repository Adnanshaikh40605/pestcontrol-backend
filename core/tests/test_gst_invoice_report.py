from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from core.gst_invoice import split_gst
from core.models import Invoice, PurchaseBill


class GstInvoiceApiTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='gst_ca', password='testpass123')
        self.client.force_authenticate(user=self.user)

    def _invoice(self, **overrides):
        data = {
            'invoice_no': 'INV-0040',
            'invoice_date': '2026-04-01',
            'billed_by_name': 'Multi Pest Care LLP',
            'billed_by_address': 'Mumbai, Maharashtra, India',
            'billed_by_gst_number': '27ACEFM4002G1ZM',
            'customer_name': 'KSHETRAPAL ART JEWELLERY',
            'customer_address': 'Mumbai',
            'customer_state': 'Maharashtra',
            'customer_gst_number': '27AYTPA2835Q1ZR',
            'supply_category': 'B2B',
            'place_of_supply': 'Maharashtra',
            'sac_code': '998531',
            'gst_rate': '18.00',
            'payment_received': '5664.00',
            'bank_ifsc': 'IDFB0040115',
            'items': [
                {
                    'service': 'Bed Bug',
                    'quantity': '2.00',
                    'rate': '2400.00',
                    'sac_code': '998531',
                    'amount': '0.00',
                }
            ],
        }
        data.update(overrides)
        return data

    def test_sample_invoice_splits_cgst_and_sgst(self):
        split = split_gst('4800.00', '18.00', 'Maharashtra', '27ACEFM4002G1ZM')
        self.assertEqual(split['cgst'], Decimal('432.00'))
        self.assertEqual(split['sgst'], Decimal('432.00'))
        self.assertEqual(split['igst'], Decimal('0.00'))
        self.assertEqual(split['grand_total'], Decimal('5664.00'))

        response = self.client.post('/api/v1/invoices/', self._invoice(), format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['subtotal'], '4800.00')
        self.assertEqual(response.data['cgst_amount'], '432.00')
        self.assertEqual(response.data['sgst_amount'], '432.00')
        self.assertEqual(response.data['grand_total'], '5664.00')
        self.assertEqual(response.data['balance_due'], '0.00')
        self.assertEqual(response.data['bank_ifsc'], 'IDFB0040115')

    def test_b2b_requires_gstin_and_unique_invoice_number(self):
        missing = self.client.post(
            '/api/v1/invoices/',
            self._invoice(customer_gst_number=''),
            format='json',
        )
        self.assertEqual(missing.status_code, status.HTTP_400_BAD_REQUEST)

        self.client.post('/api/v1/invoices/', self._invoice(), format='json')
        duplicate = self.client.post('/api/v1/invoices/', self._invoice(invoice_no='inv-0040'), format='json')
        self.assertEqual(duplicate.status_code, status.HTTP_400_BAD_REQUEST)

    def test_b2c_clears_gstin_and_other_state_uses_igst(self):
        response = self.client.post(
            '/api/v1/invoices/',
            self._invoice(
                invoice_no='INV-0041',
                supply_category='B2C',
                customer_gst_number='27AYTPA2835Q1ZR',
                place_of_supply='Karnataka',
            ),
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['customer_gst_number'], '')
        self.assertEqual(response.data['igst_amount'], '864.00')
        self.assertEqual(response.data['cgst_amount'], '0.00')

    def test_summary_deducts_only_ca_eligible_input(self):
        sales = self.client.post(
            '/api/v1/invoices/',
            self._invoice(
                invoice_no='INV-18000',
                invoice_date='2026-04-15',
                items=[{'service': 'AMC', 'amount': '100000.00', 'sac_code': '998531'}],
                payment_received='0.00',
            ),
            format='json',
        )
        self.assertEqual(sales.data['tax_amount'], '18000.00')

        cancelled = self.client.post(
            '/api/v1/invoices/',
            self._invoice(
                invoice_no='INV-CANCEL',
                invoice_date='2026-04-16',
                is_cancelled=True,
                items=[{'service': 'Cancelled visit', 'amount': '1000.00'}],
            ),
            format='json',
        )
        self.assertEqual(cancelled.status_code, status.HTTP_201_CREATED)

        PurchaseBill.objects.create(
            supplier_name='Chemical Supplier',
            supplier_gstin='27AAAAA0000A1Z5',
            bill_number='CHEM-1',
            bill_date='2026-04-10',
            taxable_amount=Decimal('20000.00'),
            igst_amount=Decimal('3600.00'),
            total_amount=Decimal('23600.00'),
            input_eligibility='eligible',
        )
        PurchaseBill.objects.create(
            supplier_name='Pending Supplier',
            supplier_gstin='27BBBBB0000B1Z5',
            bill_number='CHEM-2',
            bill_date='2026-04-11',
            taxable_amount=Decimal('1000.00'),
            igst_amount=Decimal('180.00'),
            total_amount=Decimal('1180.00'),
            input_eligibility='pending',
        )

        report = self.client.get('/api/v1/gst-reports/?month=2026-04')
        self.assertEqual(report.status_code, status.HTTP_200_OK)
        book = load_workbook(BytesIO(report.content))
        self.assertEqual(book.sheetnames, ['Sales Output', 'Purchases Input', 'Summary'])
        summary = {row[0]: row[1] for row in book['Summary'].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary['Output GST'], 18000)
        self.assertEqual(summary['Eligible input IGST (CA marked only)'], 3600)
        self.assertEqual(summary['Government payment total'], 14400)
        sales_rows = list(book['Sales Output'].iter_rows(values_only=True))
        self.assertIn('Cancelled', [row[2] for row in sales_rows])
        self.assertIn('INV-18000', [row[0] for row in sales_rows])

        settings = self.client.patch('/api/v1/gst-ca-settings/', {'ca_email': 'ca@example.com'}, format='json')
        self.assertEqual(settings.status_code, status.HTTP_200_OK)
        self.assertEqual(Invoice.objects.filter(invoice_no='INV-18000').exists(), True)
