"""GST tax split, monthly CA workbook, and the emailed package."""

from __future__ import annotations

import calendar
import io
import zipfile
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from django.db.models import Q
from openpyxl import Workbook

from core.payment_utils import quantize_money

DEFAULT_SAC = '998531'
COMPANY_STATE = 'Maharashtra'
COMPANY_STATE_CODE = '27'


def _money(value) -> Decimal:
    return quantize_money(value or 0)


def is_intra_state(place_of_supply: str, seller_gstin: str = '') -> bool:
    place = (place_of_supply or '').strip().lower()
    if 'maharashtra' in place or place in {'mh', '27'}:
        return True
    gstin = (seller_gstin or '').strip().upper()
    if place == '' and gstin.startswith(COMPANY_STATE_CODE):
        return True
    return False


def split_gst(taxable, gst_rate, place_of_supply: str, seller_gstin: str = '') -> dict:
    """Split output GST into CGST+SGST (same state) or IGST (other state)."""
    base = _money(taxable)
    rate = Decimal(str(gst_rate or 0))
    total_gst = _money(base * rate / Decimal('100'))
    if is_intra_state(place_of_supply, seller_gstin):
        half = _money(base * (rate / Decimal('2')) / Decimal('100'))
        cgst, sgst, igst = half, _money(total_gst - half), Decimal('0.00')
    else:
        cgst = sgst = Decimal('0.00')
        igst = total_gst
    return {
        'taxable': base,
        'gst_rate': rate.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
        'cgst': cgst,
        'sgst': sgst,
        'igst': igst,
        'total_gst': _money(cgst + sgst + igst),
        'grand_total': _money(base + total_gst),
    }


def month_bounds(year: int, month: int) -> tuple[date, date]:
    last = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last)


def signed_amount(invoice, value) -> Decimal:
    """Credit notes reduce output; cancelled invoices are excluded from totals."""
    amount = _money(value)
    if invoice.is_cancelled:
        return Decimal('0.00')
    if invoice.note_kind == 'credit':
        return -amount
    return amount


def document_label(invoice) -> str:
    if invoice.is_cancelled:
        return 'Cancelled'
    if invoice.note_kind == 'credit':
        return 'Credit note'
    if invoice.note_kind == 'debit':
        return 'Debit note'
    return 'Invoice'


def build_monthly_workbook(year: int, month: int) -> bytes:
    from core.models import Invoice, PurchaseBill

    start, end = month_bounds(year, month)
    invoices = (
        Invoice.objects.filter(invoice_date__gte=start, invoice_date__lte=end)
        .prefetch_related('items')
        .order_by('invoice_date', 'invoice_no')
    )
    bills = PurchaseBill.objects.filter(bill_date__gte=start, bill_date__lte=end).order_by('bill_date', 'bill_number')

    wb = Workbook()
    sales = wb.active
    sales.title = 'Sales Output'
    sales.append([
        'Invoice no.', 'Date', 'Document', 'B2B/B2C', 'Party name', 'GSTIN',
        'Location / state', 'Place of Supply', 'SAC', 'Taxable amount',
        'CGST', 'SGST', 'IGST', 'Total', 'Payment received', 'Balance due',
    ])
    out_taxable = out_cgst = out_sgst = out_igst = out_total = Decimal('0.00')
    for invoice in invoices:
        party_gstin = invoice.customer_gst_number or (
            'B2C – Unregistered' if invoice.supply_category == 'B2C' else ''
        )
        balance = _money(invoice.grand_total) - _money(invoice.payment_received)
        sales.append([
            invoice.invoice_no,
            invoice.invoice_date.isoformat(),
            document_label(invoice),
            invoice.supply_category or '',
            invoice.customer_name,
            party_gstin,
            invoice.customer_state or invoice.customer_address,
            invoice.place_of_supply,
            invoice.sac_code or DEFAULT_SAC,
            float(invoice.subtotal or 0),
            float(invoice.cgst_amount or 0),
            float(invoice.sgst_amount or 0),
            float(invoice.igst_amount or 0),
            float(invoice.grand_total or 0),
            float(invoice.payment_received or 0),
            float(balance),
        ])
        out_taxable += signed_amount(invoice, invoice.subtotal)
        out_cgst += signed_amount(invoice, invoice.cgst_amount)
        out_sgst += signed_amount(invoice, invoice.sgst_amount)
        out_igst += signed_amount(invoice, invoice.igst_amount)
        out_total += signed_amount(invoice, invoice.grand_total)

    purchases = wb.create_sheet('Purchases Input')
    purchases.append([
        'Supplier name', 'GSTIN', 'Bill no.', 'Date', 'Taxable amount',
        'CGST', 'SGST', 'IGST', 'Total', 'Bill attachment', 'Input eligibility',
    ])
    in_taxable = in_cgst = in_sgst = in_igst = in_total = Decimal('0.00')
    eligible_cgst = eligible_sgst = eligible_igst = Decimal('0.00')
    for bill in bills:
        attachment = bill.attachment.name if bill.attachment else ''
        purchases.append([
            bill.supplier_name,
            bill.supplier_gstin,
            bill.bill_number,
            bill.bill_date.isoformat(),
            float(bill.taxable_amount or 0),
            float(bill.cgst_amount or 0),
            float(bill.sgst_amount or 0),
            float(bill.igst_amount or 0),
            float(bill.total_amount or 0),
            attachment,
            bill.get_input_eligibility_display(),
        ])
        in_taxable += _money(bill.taxable_amount)
        in_cgst += _money(bill.cgst_amount)
        in_sgst += _money(bill.sgst_amount)
        in_igst += _money(bill.igst_amount)
        in_total += _money(bill.total_amount)
        if bill.input_eligibility == 'eligible':
            eligible_cgst += _money(bill.cgst_amount)
            eligible_sgst += _money(bill.sgst_amount)
            eligible_igst += _money(bill.igst_amount)

    summary = wb.create_sheet('Summary')
    payable_cgst = out_cgst - eligible_cgst
    payable_sgst = out_sgst - eligible_sgst
    payable_igst = out_igst - eligible_igst
    summary.append(['Period', f'{start.isoformat()} to {end.isoformat()}'])
    summary.append(['Sales taxable (paid and unpaid, after credit notes, excluding cancelled)', float(out_taxable)])
    summary.append(['Output CGST', float(out_cgst)])
    summary.append(['Output SGST', float(out_sgst)])
    summary.append(['Output IGST', float(out_igst)])
    summary.append(['Output GST', float(out_cgst + out_sgst + out_igst)])
    summary.append(['Sales total', float(out_total)])
    summary.append([])
    summary.append(['Purchase taxable (recorded, not automatically claimable)', float(in_taxable)])
    summary.append(['Purchase CGST recorded', float(in_cgst)])
    summary.append(['Purchase SGST recorded', float(in_sgst)])
    summary.append(['Purchase IGST recorded', float(in_igst)])
    summary.append(['Purchase total recorded', float(in_total)])
    summary.append([])
    summary.append(['Eligible input CGST (CA marked only)', float(eligible_cgst)])
    summary.append(['Eligible input SGST (CA marked only)', float(eligible_sgst)])
    summary.append(['Eligible input IGST (CA marked only)', float(eligible_igst)])
    summary.append(['Government payment CGST (output minus eligible input)', float(payable_cgst)])
    summary.append(['Government payment SGST (output minus eligible input)', float(payable_sgst)])
    summary.append(['Government payment IGST (output minus eligible input)', float(payable_igst)])
    summary.append(['Government payment total', float(payable_cgst + payable_sgst + payable_igst)])
    summary.append([])
    summary.append(['Note', 'Pending and not-eligible purchase GST is listed for GSTR-2B matching and is not deducted here.'])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_escape(text: str) -> str:
    return (
        text.replace('\\', '\\\\')
        .replace('(', '\\(')
        .replace(')', '\\)')
        .encode('latin-1', 'replace')
        .decode('latin-1')
    )


def simple_pdf(lines: list[str]) -> bytes:
    """One-page Helvetica PDF. Used when a browser is not rendering the invoice."""
    commands = ['BT', '/F1 11 Tf', '50 800 Td', '14 TL']
    for line in lines:
        commands.append(f'({_pdf_escape(line[:110])}) Tj')
        commands.append('T*')
    commands.append('ET')
    stream = '\n'.join(commands).encode('latin-1', 'replace')
    objects = []
    objects.append(b'1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n')
    objects.append(b'2 0 obj<< /Type /Pages /Count 1 /Kids [3 0 R] >>endobj\n')
    objects.append(
        b'3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] '
        b'/Contents 4 0 R /Resources<< /Font<< /F1 5 0 R >> >> >>endobj\n'
    )
    objects.append(b'4 0 obj<< /Length ' + str(len(stream)).encode() + b' >>stream\n' + stream + b'\nendstream\nendobj\n')
    objects.append(b'5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n')
    content = b'%PDF-1.4\n'
    offsets = [0]
    for obj in objects:
        offsets.append(len(content))
        content += obj
    xref = len(content)
    content += f'xref\n0 {len(offsets)}\n'.encode()
    content += b'0000000000 65535 f \n'
    for offset in offsets[1:]:
        content += f'{offset:010d} 00000 n \n'.encode()
    content += f'trailer<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
    return content


def invoice_pdf_lines(invoice) -> list[str]:
    from core.payment_utils import quantize_money as money

    balance = money(invoice.grand_total) - money(invoice.payment_received)
    party_gstin = invoice.customer_gst_number or 'B2C - Unregistered'
    lines = [
        'TAX INVOICE',
        invoice.billed_by_name or 'Multi Pest Care LLP',
        (invoice.billed_by_address or '').replace('\n', ', ')[:110],
        f'GSTIN {invoice.billed_by_gst_number}',
        f'Invoice no. {invoice.invoice_no}',
        f'Date {invoice.invoice_date.isoformat()}  Due {invoice.due_date or ""}',
        f'Terms {invoice.payment_terms}',
        f'Place of Supply {invoice.place_of_supply}',
        f'Bill to {invoice.customer_name}',
        party_gstin if invoice.supply_category != 'B2C' else 'B2C - Unregistered',
        (invoice.customer_address or '').replace('\n', ', ')[:110],
        f'State {invoice.customer_state}',
        '',
        'Description                     SAC      Amount',
    ]
    for item in invoice.items.all():
        sac = item.sac_code or invoice.sac_code or DEFAULT_SAC
        lines.append(f'{item.service[:28]:<28} {sac:<8} {item.amount}')
    lines.extend([
        '',
        f'Taxable {invoice.subtotal}',
        f'CGST {invoice.cgst_amount}   SGST {invoice.sgst_amount}   IGST {invoice.igst_amount}',
        f'Total GST {invoice.tax_amount}',
        f'Grand total {invoice.grand_total}',
        f'Payment received {invoice.payment_received}',
        f'Balance due {balance}',
        f'Bank IFSC {invoice.bank_ifsc}',
        document_label(invoice),
    ])
    return lines


def build_monthly_package(year: int, month: int) -> bytes:
    from core.models import Invoice, PurchaseBill

    start, end = month_bounds(year, month)

    archive = io.BytesIO()
    stamp = f'{year:04d}-{month:02d}'
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr(f'gst-report-{stamp}.xlsx', build_monthly_workbook(year, month))
        invoices = Invoice.objects.filter(
            invoice_date__gte=start, invoice_date__lte=end,
        ).prefetch_related('items')
        for invoice in invoices:
            bundle.writestr(
                f'sales-invoices/{invoice.invoice_no}.pdf',
                simple_pdf(invoice_pdf_lines(invoice)),
            )
        bills = PurchaseBill.objects.filter(bill_date__gte=start, bill_date__lte=end)
        for bill in bills:
            if not bill.attachment:
                continue
            try:
                bill.attachment.open('rb')
                payload = bill.attachment.read()
            except OSError:
                continue
            finally:
                bill.attachment.close()
            name = bill.attachment.name.rsplit('/', 1)[-1]
            bundle.writestr(f'purchase-bills/{bill.bill_number}-{name}', payload)
    return archive.getvalue()


def invoice_in_period(year: int, month: int):
    from core.models import Invoice

    start, end = month_bounds(year, month)
    return Invoice.objects.filter(
        Q(invoice_date__gte=start) & Q(invoice_date__lte=end),
    )
